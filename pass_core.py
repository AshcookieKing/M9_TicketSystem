"""
Ядро автоматизации заявок ММТС-9:
парсинг писем, профили компаний/гендиректоров, импорт Excel, PDF с печатью.
"""

from __future__ import annotations

import json
import re
import shutil
import sys
import uuid
from copy import deepcopy
from datetime import datetime, timedelta
from pathlib import Path


def _runtime_paths():
    """Пути для исходников (EXE) и пользовательских данных рядом с программой."""
    if getattr(sys, "frozen", False):
        bundle = Path(getattr(sys, "_MEIPASS", Path(sys.executable).parent))
        user = Path(sys.executable).resolve().parent
    else:
        bundle = Path(__file__).resolve().parent
        user = bundle
    return bundle, user


BUNDLE_DIR, APP_DIR = _runtime_paths()
DATA_DIR = APP_DIR / "data"
STAMPS_DIR = DATA_DIR / "stamps"
OUTPUT_DIR = DATA_DIR / "output"
STORE_FILE = DATA_DIR / "workspace.json"
CATALOG_FILE = DATA_DIR / "catalog.json"
BUNDLED_CATALOG = BUNDLE_DIR / "data" / "catalog.json"
TURNSTILE_ROOM = "Линия входных турникетов"

RU_MONTHS = {
    "январ": 1, "феврал": 2, "март": 3, "апрел": 4, "ма": 5,
    "июн": 6, "июл": 7, "август": 8, "сентябр": 9, "октябр": 10,
    "ноябр": 11, "декабр": 12,
}

DEFAULT_VISITORS = [
    {"id": "yugov", "surname": "Югов", "name": "Антон", "patronymic": "Дмитриевич", "birth_date": "27.09.1990", "citizenship": "Российская Федерация", "passport_series": "4510", "passport_number": "943670", "passport_issue_date": "21.10.2010", "passport_issued_by": "Отделением по району Богородское ОУФМС России по гор. Москве в ВАО", "electrical_safety_group": "III", "electrical_safety_certificate": "", "note": "", "tags": ["technocenter", "techru"]},
    {"id": "gapizov", "surname": "Гапизов", "name": "Лимат", "patronymic": "Набигуллаевич", "birth_date": "29.03.1996", "citizenship": "Российская Федерация", "passport_series": "4623", "passport_number": "378830", "passport_issue_date": "08.05.2024", "passport_issued_by": "ГУ МВД РОССИИ ПО МОСКОВСКОЙ ОБЛАСТИ", "electrical_safety_group": "III", "electrical_safety_certificate": "", "note": "", "tags": ["technocenter", "techru"]},
    {"id": "vakilov", "surname": "Вакилов", "name": "Руслан", "patronymic": "Хадимович", "birth_date": "23.05.1986", "citizenship": "Российская Федерация", "passport_series": "4611", "passport_number": "238515", "passport_issue_date": "17.12.2010", "passport_issued_by": "ТП №2 в гор. Мытищи ОУФМС России по Московаской обл. в Мытищенском р-не.", "electrical_safety_group": "III", "electrical_safety_certificate": "", "note": "", "tags": ["technocenter", "techru"]},
    {"id": "ismailov", "surname": "Исмаилов", "name": "Рустам", "patronymic": "Магомедович", "birth_date": "25.02.1997", "citizenship": "Российская Федерация", "passport_series": "4615", "passport_number": "192134", "passport_issue_date": "23.03.2017", "passport_issued_by": "ОУФМС РОССИИ ПОМОСКОВСКОЙ ОБЛ. ПО ВОЛОКОЛАМСКОМУ МУНИЦИПАЛЬНОМУ РАЙОНУ", "electrical_safety_group": "III", "electrical_safety_certificate": "", "note": "", "tags": ["technocenter", "techru"]},
    {"id": "petrov", "surname": "Петров", "name": "Евгений", "patronymic": "Сергеевич", "birth_date": "23.07.1988", "citizenship": "Российская Федерация", "passport_series": "4521", "passport_number": "644149", "passport_issue_date": "18.12.2021", "passport_issued_by": "ГУ МВД РОССИИ ПО Г.МОСКВЕ", "electrical_safety_group": "III", "electrical_safety_certificate": "4540", "note": "", "tags": ["technocenter"]},
    {"id": "rogozin", "surname": "Рогозин", "name": "Андрей", "patronymic": "Александрович", "birth_date": "18.02.1980", "citizenship": "Российская Федерация", "passport_series": "4626", "passport_number": "130612", "passport_issue_date": "25.07.2026", "passport_issued_by": "ГУ МВД России по Московской области", "electrical_safety_group": "III", "electrical_safety_certificate": "", "note": "", "tags": ["technocenter", "techru"]},
    {"id": "matyukov", "surname": "Матюков", "name": "Павел", "patronymic": "Игоревич", "birth_date": "13.09.2001", "citizenship": "Российская Федерация", "passport_series": "4521", "passport_number": "596532", "passport_issue_date": "04.11.2021", "passport_issued_by": "ГУ МВД РОССИИ ПО Г МОСКВЕ", "electrical_safety_group": "III", "electrical_safety_certificate": "", "note": "", "tags": ["technocenter", "techru"]},
    {"id": "pronyakin", "surname": "Пронякин", "name": "Алексей", "patronymic": "Владимирович", "birth_date": "08.05.1985", "citizenship": "Российская Федерация", "passport_series": "4507", "passport_number": "952490", "passport_issue_date": "07.06.2005", "passport_issued_by": "Паспортно-визовое отделение ОВД р-на Строгино гор. Москвы", "electrical_safety_group": "III", "electrical_safety_certificate": "", "note": "", "tags": ["technocenter", "techru"]},
    {"id": "ismanov", "surname": "Исманов", "name": "Мурад", "patronymic": "Акрамжонович", "birth_date": "27.10.1980", "citizenship": "Российская Федерация", "passport_series": "4625", "passport_number": "322177", "passport_issue_date": "18.11.2025", "passport_issued_by": "ГУ МВД России Московской области", "electrical_safety_group": "III", "electrical_safety_certificate": "", "note": "", "tags": ["technocenter", "techru"]},
    {"id": "ulakov", "surname": "Улаков", "name": "Дмитрий", "patronymic": "Андреевич", "birth_date": "04.10.1989", "citizenship": "Российская Федерация", "passport_series": "4512", "passport_number": "966039", "passport_issue_date": "25.01.2013", "passport_issued_by": "Отделом УФМС России по гор. Москве по району Царицыно", "electrical_safety_group": "3", "electrical_safety_certificate": "2/26", "note": "", "tags": ["techru"]},
    {"id": "mamutkin", "surname": "Мамуткин", "name": "Юрий", "patronymic": "Валерьевич", "birth_date": "25.04.1985", "citizenship": "Российская Федерация", "passport_series": "9704", "passport_number": "355605", "passport_issue_date": "09.08.2005", "passport_issued_by": "ОВД Московского р-на г. Чебоксары", "electrical_safety_group": "3", "electrical_safety_certificate": "3/26", "note": "", "tags": ["techru"]},
    {"id": "morozov", "surname": "Морозов", "name": "Виталий", "patronymic": "Вячеславович", "birth_date": "24.10.1984", "citizenship": "Российская Федерация", "passport_series": "4508", "passport_number": "161344", "passport_issue_date": "10.08.2005", "passport_issued_by": "ОВД Марьинский парк г. Москвы", "electrical_safety_group": "3", "electrical_safety_certificate": "1/26", "note": "", "tags": ["techru"]},
]


def _now_id(prefix="id"):
    return f"{prefix}_{uuid.uuid4().hex[:8]}"


def default_workspace():
    techru_director = {
        "id": "dir_krayushkin",
        "title": "Генеральный директор",
        "name": "Краюшкин С.А.",
        "stamp_file": "texru.png",
    }
    techno_director = {
        "id": "dir_savina",
        "title": "Генеральный директор",
        "name": "Савина Г.Г.",
        "stamp_file": "savina.png",
    }
    return {
        "active_company_id": "technocenter",
        "companies": [
            {
                "id": "technocenter",
                "short_name": "Техноцентр",
                "name": "ООО Техноцентр",
                "inn": "7727340108",
                "contract_number": "1319",
                "contract_date": "13.08.2021",
                "contract_valid_until": "",
                "phone": "84955892681",
                "phone_ext": "1",
                "request_prefix": "04-08",
                "last_request_seq": 823,
                "purpose_default": "Работа с оборудованием",
                "active_director_id": "dir_savina",
                "directors": [deepcopy(techno_director), deepcopy(techru_director)],
                "default_visitor_ids": [
                    "yugov", "gapizov", "vakilov", "ismailov", "petrov",
                    "rogozin", "matyukov", "pronyakin", "ismanov",
                ],
                "client_aliases": ["эйбиси", "abc-labs", "abc labs", "abclabs", "югов"],
            },
            {
                "id": "techru",
                "short_name": "Тех РУ",
                "name": "ООО Тех РУ",
                "inn": "9717010270",
                "contract_number": "105",
                "contract_date": "01.10.2010",
                "contract_valid_until": "01.02.2025",
                "phone": "84955892681",
                "phone_ext": "1",
                "request_prefix": "04-08",
                "last_request_seq": 825,
                "purpose_default": "Работа с оборудованием",
                "active_director_id": "dir_krayushkin",
                "directors": [deepcopy(techru_director), deepcopy(techno_director)],
                "default_visitor_ids": [
                    "pronyakin", "ismanov", "ulakov", "mamutkin", "gapizov",
                    "vakilov", "ismailov", "morozov", "rogozin", "matyukov",
                ],
                "client_aliases": ["нца", "тех ру", "тех.ру", "tech.ru"],
            },
        ],
        "visitors": deepcopy(DEFAULT_VISITORS),
        "clients": [
            {"name": "ABC-Labs / ЭЙБИСИ-ЛАБС", "aliases": ["эйбиси", "abc-labs", "abc labs"], "company_id": "technocenter"},
            {"name": "НЦА", "aliases": ["нца"], "company_id": "techru"},
        ],
    }


def ensure_dirs():
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    STAMPS_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    bundled = BUNDLE_DIR / "data" / "stamps"
    if bundled.exists() and bundled.resolve() != STAMPS_DIR.resolve():
        for src in bundled.iterdir():
            dest = STAMPS_DIR / src.name
            if src.is_file() and not dest.exists():
                shutil.copy2(src, dest)


def _seed_stamps(data):
    """Подставляет печати в уже сохранённый профиль."""
    for company in data.get("companies", []):
        for director in company.get("directors", []):
            if director.get("id") == "dir_savina" and not director.get("stamp_file"):
                if (STAMPS_DIR / "savina.png").exists():
                    director["stamp_file"] = "savina.png"
            if director.get("id") == "dir_krayushkin":
                if (STAMPS_DIR / "texru.png").exists():
                    director["stamp_file"] = "texru.png"
                elif not director.get("stamp_file") and (STAMPS_DIR / "techru.png").exists():
                    director["stamp_file"] = "techru.png"
    return data


def _short_signer_name(person):
    surname = (person.get("surname") or "").strip()
    name = (person.get("name") or "").strip()
    patronymic = (person.get("patronymic") or "").strip()
    initials = "".join(f"{part[0]}." for part in (name, patronymic) if part)
    return f"{surname} {initials}".strip()


def default_line_signers():
    items = [
        {"title": "Инженер", "name": "Порозов С.Ю", "email": "m9@tech.ru"},
        {"title": "Вед.инж", "name": "Пронякин А.В.", "email": "support@tech.ru"},
    ]
    seen_names = {item["name"] for item in items}
    seen_surnames = {item["name"].split()[0] for item in items if item.get("name")}
    for person in DEFAULT_VISITORS:
        if "techru" not in (person.get("tags") or []):
            continue
        name = _short_signer_name(person)
        surname = (person.get("surname") or "").strip()
        if not name or name in seen_names or surname in seen_surnames:
            continue
        seen_names.add(name)
        seen_surnames.add(surname)
        items.append({"title": "Инженер", "name": name, "email": "m9@tech.ru"})
    return items


class Workspace:
    def __init__(self):
        ensure_dirs()
        self.data = self.load()
        self._ensure_line_signers()

    def _ensure_line_signers(self):
        changed = False
        defaults = default_line_signers()
        for company in self.companies():
            signers = company.get("line_signers") or []
            known = {(item.get("name") or "").strip() for item in signers}
            if not signers:
                company["line_signers"] = deepcopy(defaults)
                changed = True
                continue
            for item in defaults:
                name = (item.get("name") or "").strip()
                if name and name not in known:
                    signers.append(deepcopy(item))
                    known.add(name)
                    changed = True
            company["line_signers"] = signers
        if changed:
            self.save()

    def upsert_line_signer(self, name, title="Инженер", email=""):
        company = self.get_company()
        if not company:
            return None
        name = (name or "").strip()
        if not name:
            return None
        signers = company.setdefault("line_signers", [])
        for signer in signers:
            if (signer.get("name") or "").strip() == name:
                if title:
                    signer["title"] = (title or "").strip() or signer.get("title") or "Инженер"
                if email:
                    signer["email"] = (email or "").strip()
                self.save()
                return signer
        signer = {
            "title": (title or "Инженер").strip() or "Инженер",
            "name": name,
            "email": (email or "").strip(),
        }
        signers.append(signer)
        self.save()
        return signer

    def load(self):
        if STORE_FILE.exists():
            try:
                data = json.loads(STORE_FILE.read_text(encoding="utf-8"))
                return _seed_stamps(data)
            except Exception:
                pass
        data = default_workspace()
        self._save(data)
        return data

    def _save(self, data=None):
        payload = data if data is not None else self.data
        STORE_FILE.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    def save(self):
        self._save()

    def companies(self):
        return self.data.setdefault("companies", [])

    def visitors(self):
        return self.data.setdefault("visitors", [])

    def get_company(self, company_id=None):
        cid = company_id or self.data.get("active_company_id")
        for company in self.companies():
            if company["id"] == cid:
                return company
        return self.companies()[0] if self.companies() else None

    def set_active_company(self, company_id):
        self.data["active_company_id"] = company_id
        self.save()

    def get_director(self, company=None):
        company = company or self.get_company()
        if not company:
            return None
        did = company.get("active_director_id")
        for director in company.get("directors", []):
            if director["id"] == did:
                return director
        directors = company.get("directors") or []
        return directors[0] if directors else None

    def set_active_director(self, director_id, company_id=None):
        company = self.get_company(company_id)
        if not company:
            return
        company["active_director_id"] = director_id
        self.save()

    def stamp_path(self, director=None):
        director = director or self.get_director()
        if not director:
            return None
        filename = (director.get("stamp_file") or "").strip()
        if not filename:
            return None
        path = STAMPS_DIR / Path(filename).name
        if path.exists():
            return path
        legacy = APP_DIR / filename
        if legacy.exists():
            return legacy
        return None

    def save_stamp(self, source_path, director):
        source = Path(source_path)
        ext = source.suffix.lower() or ".png"
        filename = f"{director['id']}{ext}"
        dest = STAMPS_DIR / filename
        shutil.copyfile(source, dest)
        director["stamp_file"] = filename
        self.save()
        return dest

    def next_request_number(self, company=None, consume=False):
        company = company or self.get_company()
        seq = int(company.get("last_request_seq") or 0) + 1
        prefix = company.get("request_prefix") or "04-08"
        number = f"{prefix}.{seq}"
        if consume:
            company["last_request_seq"] = seq
            self.save()
        return number

    def peek_request_number(self, company=None):
        return self.next_request_number(company, consume=False)

    def consume_request_number(self, company=None, number=None):
        company = company or self.get_company()
        if number:
            match = re.search(r"(\d+)$", str(number))
            if match:
                company["last_request_seq"] = max(int(company.get("last_request_seq") or 0), int(match.group(1)))
                self.save()
                return number
        return self.next_request_number(company, consume=True)

    def visitor_by_id(self, visitor_id):
        for visitor in self.visitors():
            if visitor.get("id") == visitor_id:
                return visitor
        return None

    def default_team(self, company=None):
        company = company or self.get_company()
        team = []
        for vid in company.get("default_visitor_ids", []):
            visitor = self.visitor_by_id(vid)
            if visitor:
                team.append(visitor_to_request(visitor, company.get("name", "")))
        return team

    def find_visitors_in_text(self, text, company=None):
        found = []
        seen = set()
        lowered = (text or "").lower()
        for visitor in self.visitors():
            key = visitor["id"]
            if key in seen:
                continue
            if visitor["surname"].lower() in lowered:
                found.append(visitor_to_request(visitor, (company or {}).get("name", "")))
                seen.add(key)
        return found

    def suggest_company_from_text(self, text):
        lowered = (text or "").lower()
        for client in self.data.get("clients", []):
            if any(alias in lowered for alias in client.get("aliases", [])):
                return client.get("company_id")
        for company in self.companies():
            if any(alias in lowered for alias in company.get("client_aliases", [])):
                return company["id"]
            if company.get("short_name", "").lower() in lowered or company.get("name", "").lower() in lowered:
                return company["id"]
        return None

    def upsert_visitor(self, visitor):
        if not visitor.get("id"):
            visitor["id"] = _now_id("vis")
        for idx, existing in enumerate(self.visitors()):
            if existing["id"] == visitor["id"]:
                self.visitors()[idx] = visitor
                self.save()
                return visitor
        self.visitors().append(visitor)
        self.save()
        return visitor

    def delete_visitor(self, visitor_id):
        self.data["visitors"] = [v for v in self.visitors() if v.get("id") != visitor_id]
        for company in self.companies():
            company["default_visitor_ids"] = [vid for vid in company.get("default_visitor_ids", []) if vid != visitor_id]
        self.save()

    def upsert_company(self, company):
        if not company.get("id"):
            company["id"] = _now_id("co")
        for idx, existing in enumerate(self.companies()):
            if existing["id"] == company["id"]:
                self.companies()[idx] = company
                self.save()
                return company
        self.companies().append(company)
        self.save()
        return company

    def upsert_director(self, company_id, director):
        company = self.get_company(company_id)
        if not company:
            return None
        if not director.get("id"):
            director["id"] = _now_id("dir")
        directors = company.setdefault("directors", [])
        for idx, existing in enumerate(directors):
            if existing["id"] == director["id"]:
                directors[idx] = director
                self.save()
                return director
        directors.append(director)
        if not company.get("active_director_id"):
            company["active_director_id"] = director["id"]
        self.save()
        return director

    def delete_director(self, company_id, director_id):
        company = self.get_company(company_id)
        if not company:
            return
        company["directors"] = [d for d in company.get("directors", []) if d.get("id") != director_id]
        if company.get("active_director_id") == director_id:
            company["active_director_id"] = company["directors"][0]["id"] if company["directors"] else ""
        self.save()


def visitor_to_request(visitor, organization=""):
    data = {
        "surname": visitor.get("surname", ""),
        "name": visitor.get("name", ""),
        "patronymic": visitor.get("patronymic", ""),
        "birth_date": visitor.get("birth_date", ""),
        "citizenship": visitor.get("citizenship", "Российская Федерация"),
        "passport_series": visitor.get("passport_series", ""),
        "passport_number": visitor.get("passport_number", ""),
        "passport_issue_date": visitor.get("passport_issue_date", ""),
        "passport_issued_by": visitor.get("passport_issued_by", ""),
        "organization": organization or visitor.get("organization", ""),
        "electrical_safety_group": visitor.get("electrical_safety_group", ""),
        "electrical_safety_certificate": visitor.get("electrical_safety_certificate", ""),
        "note": visitor.get("note", ""),
        "pool_id": visitor.get("id", ""),
    }
    return data


_CATALOG_CACHE = None

UNITS = ["шт.", "упак.", "кор.", "м.", "кв.м.", "куб.м.", "кг.", "т.", "паллет", "бухта"]
ACTIONS = ["Внос", "Вынос"]
PURPOSES = [
    "Работа с оборудованием",
    "ТО",
    "установка дополнительного оборудования",
    "монтаж",
    "демонтаж",
]
OBJECT_ADDRESSES = [
    "МОСКВА УЛ. БУТЛЕРОВА, 7",
    "НИЖНИЙ НОВГОРОД УЛ. ФЕДОСЕЕНКО, 35",
]
SERIAL_REASONS = [
    "Отсутствует на корпусе",
    "Не читаем / поврежден",
    "Не подлежит указанию для данного вида (типа) ТМЦ",
]


def _sheet_column(ws, col, start=2):
    values = []
    seen = set()
    for row in range(start, (ws.max_row or 1) + 1):
        raw = ws.cell(row, col).value
        if raw in (None, ""):
            continue
        text = str(raw).strip()
        if not text or text in seen:
            continue
        seen.add(text)
        values.append(text)
    return values


def load_catalog():
    """Справочники с листа «Список». Сначала кэш JSON — Excel при старте не читаем."""
    global _CATALOG_CACHE
    if _CATALOG_CACHE is not None:
        return _CATALOG_CACHE
    catalog = {
        "rooms_butlerova": [TURNSTILE_ROOM],
        "rooms_fedoseenko": ["Уличная территория Федосеенко"],
        "units": list(UNITS),
        "reasons": list(SERIAL_REASONS),
        "addresses": list(OBJECT_ADDRESSES),
        "rows": [],
        "places": [],
        "actions": list(ACTIONS),
        "purposes": list(PURPOSES),
    }
    for candidate in (CATALOG_FILE, BUNDLED_CATALOG):
        if candidate.exists():
            try:
                cached = json.loads(candidate.read_text(encoding="utf-8"))
                if cached.get("rooms_butlerova"):
                    catalog.update(cached)
                    _CATALOG_CACHE = catalog
                    return catalog
            except Exception:
                pass
    try:
        from pass_automation_v2 import PassRequestAutomation
        from openpyxl import load_workbook
        automation = PassRequestAutomation()
        wb = load_workbook(automation.template_file, data_only=True, read_only=True)
        try:
            if "Список" in wb.sheetnames:
                ws = wb["Список"]
                butlerova = _sheet_column(ws, 1)
                fedoseenko = _sheet_column(ws, 14)
                units = _sheet_column(ws, 5)
                reasons = _sheet_column(ws, 16)
                addresses = _sheet_column(ws, 12)
                rows = _sheet_column(ws, 20)
                places = _sheet_column(ws, 22)
                if butlerova:
                    catalog["rooms_butlerova"] = butlerova
                if fedoseenko:
                    catalog["rooms_fedoseenko"] = fedoseenko
                if units:
                    catalog["units"] = units
                if reasons:
                    catalog["reasons"] = reasons
                if addresses:
                    catalog["addresses"] = addresses
                catalog["rows"] = rows
                catalog["places"] = places
        finally:
            wb.close()
        ensure_dirs()
        CATALOG_FILE.write_text(json.dumps(catalog, ensure_ascii=False), encoding="utf-8")
    except Exception:
        pass
    if TURNSTILE_ROOM not in catalog["rooms_butlerova"]:
        catalog["rooms_butlerova"].insert(0, TURNSTILE_ROOM)
    _CATALOG_CACHE = catalog
    return catalog


def load_rooms(address=None):
    catalog = load_catalog()
    if address and "ФЕДОСЕЕНКО" in str(address).upper():
        rooms = list(catalog["rooms_fedoseenko"])
    else:
        rooms = list(catalog["rooms_butlerova"])
    return rooms


def default_access_room(address=None):
    rooms = load_rooms(address)
    return rooms[0] if rooms else TURNSTILE_ROOM


class RoomCatalog:
    def __init__(self, rooms=None):
        self.rooms = rooms or load_rooms()

    def match(self, token):
        raw = (token or "").strip()
        if not raw:
            return ""
        compact = re.sub(r"\s+", "", raw.lower().replace("№", "").replace("no", ""))
        # точное вхождение
        for room in self.rooms:
            hay = re.sub(r"\s+", "", room.lower().replace("№", ""))
            if compact and compact in hay:
                if compact.startswith("11.29") and "1129а" in hay:
                    continue
                return room
        match = re.search(r"(\d+)\s*[./]\s*(\d+[а-яa-z]?)", raw, re.IGNORECASE)
        if not match:
            match = re.search(r"(\d+)\s*эт[.\s]*\w*\s*(?:бл[.\s]*\w*)?\s*(?:пом[.\s]*)?№?\s*(\d+[.\dа-яa-z]*)", raw, re.IGNORECASE)
        if not match:
            return raw
        floor, room_no = match.group(1), match.group(2).upper().replace("А", "А")
        floor_rooms = [r for r in self.rooms if re.match(rf"^{floor}\s*эт", r, re.IGNORECASE)]
        exact_dot = f"№{floor}.{room_no}".lower()
        for room in floor_rooms:
            if exact_dot in room.lower().replace(" ", ""):
                return room
        concat = f"№{floor}{room_no}".lower()
        for room in floor_rooms:
            hay = room.lower().replace(" ", "")
            if concat in hay and not any(hay.endswith(sfx) for sfx in ("а", "б", "в", "г", "д")):
                return room
        for room in floor_rooms:
            if f"№{room_no}".lower() in room.lower().replace(" ", ""):
                return room
        for room in floor_rooms:
            if concat in room.lower().replace(" ", ""):
                return room
        return raw


def normalize_date(value, default_year=None):
    if not value:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y")
    text = str(value).strip()
    if not text or text.startswith("00.01.1900"):
        return ""
    text = text.replace("/", ".").replace("-", ".")
    parts = [p for p in text.split(".") if p]
    if len(parts) == 3 and all(p.isdigit() for p in parts[:2]):
        d, m, y = parts
        if len(y) == 2:
            y = "20" + y if int(y) < 50 else "19" + y
        return f"{d.zfill(2)}.{m.zfill(2)}.{y.zfill(4)}"
    return text


def parse_human_date(text, year=None):
    year = year or datetime.now().year
    text = (text or "").lower()
    match = re.search(
        r"(\d{1,2})\s*(январ\w*|феврал\w*|март\w*|апрел\w*|ма[йя]\w*|июн\w*|июл\w*|август\w*|сентябр\w*|октябр\w*|ноябр\w*|декабр\w*)(?:\s*(\d{4}))?",
        text,
    )
    if not match:
        return ""
    day = int(match.group(1))
    month_raw = match.group(2)
    year = int(match.group(3) or year)
    month = 5 if month_raw.startswith("ма") and not month_raw.startswith("март") else None
    if month is None:
        for key, value in RU_MONTHS.items():
            if month_raw.startswith(key):
                month = value
                break
    if not month:
        return ""
    try:
        return datetime(year, month, day).strftime("%d.%m.%Y")
    except ValueError:
        return ""


def add_days(date_str, days):
    dt = datetime.strptime(date_str, "%d.%m.%Y")
    return (dt + timedelta(days=days)).strftime("%d.%m.%Y")


def format_phone(phone, ext=""):
    phone = re.sub(r"\s+", "", phone or "")
    if not phone:
        return ""
    if ext:
        return f"{phone} доб.{ext}"
    return phone


class EmailParser:
    def __init__(self, rooms=None):
        self.catalog = RoomCatalog(rooms)

    def parse(self, text):
        text = (text or "").replace("\xa0", " ")
        lines = [re.sub(r"\s{2,}", " ", line.strip(" >|\t")) for line in text.splitlines()]
        joined = "\n".join(lines)
        result = {
            "client": "",
            "client_contract": "",
            "work_start_date": "",
            "work_end_date": "",
            "purpose": "",
            "contact_phone": "",
            "sender_name": "",
            "rooms": [],
            "equipment": [],
            "notes": [],
        }

        contract = re.search(r"договор\s*№?\s*([0-9A-Za-zА-Яа-я/\-_]+)", joined, re.IGNORECASE)
        if contract:
            result["client_contract"] = contract.group(1).strip(" .;")

        client = re.search(
            r"договор[^\n]*?\s+с\s+([A-ZА-ЯЁ][A-ZА-ЯЁa-zа-яё\- \"«»]+)",
            joined,
            re.IGNORECASE,
        )
        if client:
            result["client"] = client.group(1).strip(" .;")
        if not result["client"]:
            company = re.search(r'компани[яи]\s+"?([^"\n]+)"?', joined, re.IGNORECASE)
            if company:
                result["client"] = company.group(1).strip()

        start = parse_human_date(joined)
        if not start:
            date_match = re.search(r"(\d{1,2}[./-]\d{1,2}[./-]\d{2,4})", joined)
            if date_match:
                start = normalize_date(date_match.group(1))
        if start:
            result["work_start_date"] = start
            period = re.search(
                r"(?:с|от)\s*(\d{1,2}[./-]\d{1,2}[./-]\d{2,4}|\d{1,2}\s+[а-яё]+)\s*(?:по|до)\s*(\d{1,2}[./-]\d{1,2}[./-]\d{2,4}|\d{1,2}\s+[а-яё]+)",
                joined,
                re.IGNORECASE,
            )
            if period:
                end = parse_human_date(period.group(2)) or normalize_date(period.group(2))
                result["work_end_date"] = end or add_days(start, 3)
            else:
                result["work_end_date"] = add_days(start, 3)

        if re.search(r"\bто\b|техобслуж|обслуживан", joined, re.IGNORECASE):
            result["purpose"] = "ТО"
        elif re.search(r"установк", joined, re.IGNORECASE):
            result["purpose"] = "установка дополнительного оборудования"
        else:
            result["purpose"] = "Работа с оборудованием"

        phone = re.search(r"(?:тел\.?|телефон)[:\s]*([+\d\s\-()]{7,})", joined, re.IGNORECASE)
        if phone:
            result["contact_phone"] = re.sub(r"\s+", " ", phone.group(1)).strip(" -")

        sender = re.search(r"с уважением[,:\s]+([А-ЯЁ][а-яё]+(?:\s+[А-ЯЁ][а-яё]+){0,2})", joined, re.IGNORECASE)
        if sender:
            result["sender_name"] = sender.group(1).strip()

        result["equipment"] = self._parse_equipment(joined)
        rooms = []
        for item in result["equipment"]:
            if item.get("room") and item["room"] not in rooms:
                rooms.append(item["room"])
        result["rooms"] = rooms
        if "из заявки" in joined.lower() or "во вложении" in joined.lower():
            result["notes"].append("В письме указано вложение — прикрепите Excel клиента для серийников и состава.")
        return result

    def _parse_equipment(self, text):
        items = []
        location_re = re.compile(
            r"(?P<room>\d{1,2}\s*[./]\s*\d{1,3}[а-яА-Яa-zA-Z]?)"
            r"(?:\s*[.,]?\s*ряд\s*(?P<row>\d+[а-яА-Яa-zA-Z]?))?"
            r"(?:\s*место\s*(?P<place>\d+[а-яА-Яa-zA-Z]?))?",
            re.IGNORECASE,
        )
        serial_re = re.compile(r"\b(?:s/n|sn|серий(?:ный)?|зав(?:одской)?)[.\s№:]*([A-Za-z0-9\-]+)\b", re.IGNORECASE)
        bare_serial_re = re.compile(r"\b([A-Z]{0,4}\d{4,}[A-Z0-9\-]*)\b")

        action_hits = list(re.finditer(r"внос|вынос|ввоз|вывоз", text, re.IGNORECASE))
        for idx_hit, match in enumerate(action_hits):
            end = action_hits[idx_hit + 1].start() if idx_hit + 1 < len(action_hits) else len(text)
            sentence = text[match.start():min(end, match.start() + 280)]
            low = sentence.lower()
            action = "Вынос" if re.search(r"вынос|вывоз", low) else "Внос"
            name = "Сервер" if "сервер" in low else ("Жесткий диск" if re.search(r"диск|ssd|hdd", low) else "Оборудование")
            qty_word = re.search(r"(двух|три|трех|трёх|четырех|четыре|(\d+)\s*(?:шт|единиц|сервер|диск))", low)
            qty = 1
            if qty_word:
                mapping = {"двух": 2, "три": 3, "трех": 3, "трёх": 3, "четырех": 4, "четыре": 4}
                if qty_word.group(1) in mapping:
                    qty = mapping[qty_word.group(1)]
                elif qty_word.group(2):
                    qty = int(qty_word.group(2))
            serials = [m.group(1) for m in serial_re.finditer(sentence)]
            if not serials:
                for match in bare_serial_re.finditer(sentence):
                    token = match.group(1)
                    if token.isdigit() and len(token) >= 4:
                        serials.append(token)
            locations = list(location_re.finditer(sentence))
            if not locations:
                continue
            count = max(len(locations), qty, len(serials) or 1)
            for idx in range(count):
                loc = locations[idx] if idx < len(locations) else locations[-1]
                room = self.catalog.match(re.sub(r"\s+", "", loc.group("room")))
                items.append({
                    "name": name,
                    "brand": "",
                    "model": "",
                    "serial_number": serials[idx] if idx < len(serials) else "",
                    "quantity": "1",
                    "unit": "шт.",
                    "action": action,
                    "room": room,
                    "reason": "",
                    "row": (loc.group("row") or "").upper(),
                    "place": loc.group("place") or "",
                    "note": "",
                })
        return items


def _cell(row, *keys):
    for key in keys:
        if key < len(row) and row[key] not in (None, ""):
            return row[key]
    return ""


def _is_empty_name(value):
    text = str(value or "").strip()
    return (not text) or text in {"0", "None"} or text.startswith("#")


def import_client_excel(path, rooms=None):
    from openpyxl import load_workbook

    catalog = RoomCatalog(rooms)
    wb = load_workbook(path, data_only=True)
    result = {
        "request_data": {},
        "visitors": [],
        "equipment": [],
        "vehicles": [],
        "source_name": Path(path).name,
    }
    try:
        if "Данные заявки" in wb.sheetnames:
            ws = wb["Данные заявки"]
            result["request_data"] = {
                "organization": str(ws["C2"].value or "").strip(),
                "inn": str(ws["C3"].value or "").strip(),
                "contract_number": str(ws["C4"].value or "").strip(),
                "contract_date": normalize_date(ws["E4"].value),
                "contract_valid_until": normalize_date(ws["C5"].value),
                "work_description": str(ws["C7"].value or "").strip(),
                "work_start_date": normalize_date(ws["C8"].value),
                "work_end_date": normalize_date(ws["C9"].value),
                "purpose": str(ws["C10"].value or "").strip(),
                "responsible_person": str(ws["C11"].value or "").strip(),
                "responsible_name": str(ws["D11"].value or "").strip(),
                "request_number": str(ws["C12"].value or "").strip(),
                "contact_phone": str(ws["C13"].value or "").strip(),
                "object_address": str(ws["C14"].value or "").strip(),
            }
            rooms_found = []
            for col in range(3, 9):
                value = ws.cell(6, col).value
                if value:
                    rooms_found.append(str(value).strip())
            result["request_data"]["rooms"] = rooms_found
        if "Посетители" in wb.sheetnames:
            ws = wb["Посетители"]
            for row in ws.iter_rows(min_row=3, max_row=22, max_col=15, values_only=True):
                surname = str(row[2] or "").strip()
                name = str(row[3] or "").strip()
                if not surname or not name:
                    continue
                result["visitors"].append({
                    "surname": surname,
                    "name": name,
                    "patronymic": str(row[4] or "").strip(),
                    "birth_date": normalize_date(row[5]),
                    "citizenship": str(row[6] or "").strip() or "Российская Федерация",
                    "passport_series": str(row[7] or "").replace(" ", "").strip(),
                    "passport_number": str(row[8] or "").strip(),
                    "passport_issue_date": normalize_date(row[9]),
                    "passport_issued_by": str(row[10] or "").strip(),
                    "organization": str(row[11] or "").strip(),
                    "electrical_safety_group": str(row[12] or "").strip("_ "),
                    "electrical_safety_certificate": str(row[13] or "").strip("_ "),
                    "note": str(row[14] or "").strip() if len(row) > 14 and row[14] else "",
                })
        if "Оборудование" in wb.sheetnames:
            ws = wb["Оборудование"]
            for row in ws.iter_rows(min_row=3, max_row=42, max_col=16, values_only=True):
                name = row[2]
                if _is_empty_name(name):
                    continue
                raw_room = str(row[10] or "").strip()
                result["equipment"].append({
                    "name": str(name).strip(),
                    "brand": str(row[3] or "").strip(),
                    "model": str(row[4] or "").strip(),
                    "serial_number": str(row[5] or "").strip(),
                    "quantity": str(row[6] or "1").strip() or "1",
                    "unit": str(row[7] or "шт.").strip() or "шт.",
                    "action": str(row[8] or "").strip(),
                    "note": str(row[9] or "").strip(),
                    "room": catalog.match(raw_room) if raw_room else "",
                    "reason": str(row[12] or "").strip(),
                    "row": str(row[13] or "").strip(),
                    "place": str(row[14] or "").strip(),
                })
                if not result["equipment"][-1]["row"] and row[11]:
                    combo = str(row[11])
                    parts = re.split(r"\s*,\s*", combo)
                    if parts:
                        result["equipment"][-1]["row"] = parts[0].strip()
                    if len(parts) > 1:
                        result["equipment"][-1]["place"] = parts[1].strip()
        if "Автотранспорт" in wb.sheetnames:
            ws = wb["Автотранспорт"]
            for row in ws.iter_rows(min_row=4, max_row=23, max_col=14, values_only=True):
                plate = str(row[3] or "").strip()
                if not plate:
                    continue
                result["vehicles"].append({
                    "license_plate": plate,
                    "driver_surname": str(row[4] or "").strip(),
                    "driver_name": str(row[5] or "").strip(),
                    "driver_patronymic": str(row[6] or "").strip(),
                    "driver_birth_date": normalize_date(row[7]),
                    "driver_citizenship": str(row[8] or "").strip(),
                    "passport_series": str(row[9] or "").strip(),
                    "passport_number": str(row[10] or "").strip(),
                    "passport_issue_date": normalize_date(row[11]),
                    "passport_issued_by": str(row[12] or "").strip(),
                    "driver_organization": str(row[13] or "").strip() if len(row) > 13 else "",
                })
    finally:
        wb.close()
    return result


def merge_unique_visitors(*groups):
    merged = []
    seen = set()
    for group in groups:
        for visitor in group or []:
            key = (
                visitor.get("surname", "").strip().lower(),
                visitor.get("name", "").strip().lower(),
                visitor.get("passport_number", "").strip(),
            )
            if not key[0] or key in seen:
                continue
            seen.add(key)
            merged.append(visitor)
    return merged


def collect_rooms(equipment, extra=None, address=None):
    default = default_access_room(address)
    rooms = [default]
    for source in (extra or []):
        if isinstance(source, str) and source and source not in rooms:
            rooms.append(source)
        elif isinstance(source, list):
            for item in source:
                if item and item not in rooms:
                    rooms.append(item)
    for item in equipment or []:
        room = (item.get("room") or "").strip()
        if room and room not in rooms:
            rooms.append(room)
    return rooms[:6]


def build_work_description(equipment):
    parts = []
    seen = set()
    for item in equipment or []:
        room = item.get("room") or ""
        row = item.get("row") or ""
        place = item.get("place") or ""
        chunk = " ".join(filter(None, [
            room,
            f"ряд {row}" if row else "",
            f"место {place}" if place else "",
        ])).strip()
        if chunk and chunk not in seen:
            seen.add(chunk)
            parts.append(chunk)
    return "; ".join(parts)


def apply_company_defaults(request_data, company, director, request_date=None):
    phone = format_phone(company.get("phone", ""), company.get("phone_ext", ""))
    request_data.update({
        "organization": company.get("name", ""),
        "inn": company.get("inn", ""),
        "contract_number": company.get("contract_number", ""),
        "contract_date": company.get("contract_date", ""),
        "contract_valid_until": company.get("contract_valid_until", ""),
        "responsible_person": (director or {}).get("title") or "Генеральный директор",
        "responsible_name": (director or {}).get("name") or "",
        "contact_phone": phone,
        "request_date": request_date or datetime.now().strftime("%d.%m.%Y"),
        "object_address": request_data.get("object_address") or OBJECT_ADDRESSES[0],
    })
    if not request_data.get("purpose"):
        request_data["purpose"] = company.get("purpose_default") or "Работа с оборудованием"
    if not request_data.get("request_number"):
        request_data["request_number"] = ""
    return request_data


def export_excel(request_data, visitors, equipment, vehicles, output_path):
    from pass_automation_v2 import PassRequestAutomation
    automation = PassRequestAutomation()
    automation.load_template()
    automation.fill_request_data(request_data)
    automation.fill_visitors(visitors)
    automation.fill_equipment(equipment)
    automation.fill_vehicles(vehicles)
    automation.save(str(output_path))
    return output_path


def _register_pdf_fonts():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    windir = Path(os_environ_windir())
    candidates = [
        (windir / "Fonts" / "arial.ttf", windir / "Fonts" / "arialbd.ttf", "Arial", "Arial-Bold"),
        (windir / "Fonts" / "calibri.ttf", windir / "Fonts" / "calibrib.ttf", "Calibri", "Calibri-Bold"),
        (APP_DIR / "DejaVuSans.ttf", APP_DIR / "DejaVuSans-Bold.ttf", "DejaVuSans", "DejaVuSans-Bold"),
    ]
    for regular, bold, name, bold_name in candidates:
        if regular.exists():
            try:
                pdfmetrics.registerFont(TTFont(name, str(regular)))
                if bold.exists():
                    pdfmetrics.registerFont(TTFont(bold_name, str(bold)))
                else:
                    bold_name = name
                return name, bold_name
            except Exception:
                continue
    return "Helvetica", "Helvetica-Bold"


def os_environ_windir():
    import os
    return os.environ.get("WINDIR", r"C:\Windows")


def draw_round_stamp(pdf, stamp_path, page_width, signature_y, size_mm=44):
    """Круглая печать по центру блока подписи, чуть ниже строки должности — как у Краюшкина."""
    from reportlab.lib.units import mm

    if not stamp_path or not Path(stamp_path).exists():
        return
    stamp_size = size_mm * mm
    try:
        pdf.drawImage(
            str(stamp_path),
            page_width / 2 - stamp_size / 2,
            signature_y - 34 * mm,
            width=stamp_size,
            height=stamp_size,
            mask="auto",
            preserveAspectRatio=True,
            anchor="c",
        )
    except Exception:
        pass


def export_pdf(request_data, visitors, equipment, vehicles, stamp_path, output_path):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas as pdf_canvas
    from reportlab.platypus import Paragraph, Table, TableStyle

    font, font_bold = _register_pdf_fonts()
    width, height = A4
    left = 18 * mm
    right = width - 18 * mm
    pdf = pdf_canvas.Canvas(str(output_path), pagesize=A4)
    y = height - 16 * mm
    bottom = 16 * mm

    def ensure_space(need=16):
        nonlocal y
        if y < bottom + need:
            pdf.showPage()
            y = height - 16 * mm

    def draw_text(text, *, x=left, size=10, font_name=font, align="left", leading=13, max_width=None):
        nonlocal y
        max_width = max_width or (right - x)
        words = (text or "").split()
        if not words:
            y -= leading
            return
        line = ""
        pdf.setFont(font_name, size)
        for word in words:
            candidate = (line + " " + word).strip()
            if pdf.stringWidth(candidate, font_name, size) <= max_width:
                line = candidate
            else:
                ensure_space()
                pdf.setFont(font_name, size)
                if align == "center":
                    pdf.drawCentredString((left + right) / 2, y, line)
                elif align == "right":
                    pdf.drawRightString(right, y, line)
                else:
                    pdf.drawString(x, y, line)
                y -= leading
                line = word
        if line:
            ensure_space()
            pdf.setFont(font_name, size)
            if align == "center":
                pdf.drawCentredString((left + right) / 2, y, line)
            elif align == "right":
                pdf.drawRightString(right, y, line)
            else:
                pdf.drawString(x, y, line)
            y -= leading

    org = request_data.get("organization", "")
    inn = request_data.get("inn", "")
    phone = request_data.get("contact_phone", "")
    req_no = request_data.get("request_number", "")
    req_date = request_data.get("request_date", "")
    pdf.setFont(font, 10)
    pdf.drawRightString(right, height - 16 * mm, 'В АО "ММТС-9"')
    y = height - 16 * mm
    draw_text(org, size=11, font_name=font_bold)
    if inn:
        draw_text(f"ИНН {inn}")
    if phone:
        draw_text(f"тел. ответственного {phone}")
    if req_no or req_date:
        draw_text(" / ".join(part for part in (f"исх. №{req_no}" if req_no else "", req_date) if part))
    y -= 4
    obj = request_data.get("object_address") or OBJECT_ADDRESSES[0]
    draw_text("Заявка на доступ на объект", size=12, font_name=font_bold)
    draw_text(obj, size=12, font_name=font_bold)
    y -= 2
    draw_text("Заявка", size=12, font_name=font_bold, align="center")
    y -= 2

    contract_number = request_data.get("contract_number", "")
    contract_date = request_data.get("contract_date", "")
    start = request_data.get("work_start_date", "")
    end = request_data.get("work_end_date", "")
    purpose = request_data.get("purpose", "")
    paragraph = "Просим разрешить"
    if contract_number:
        paragraph += f" в соответствии с договором {contract_number}"
        if contract_date:
            paragraph += f" от {contract_date}"
    if start and end:
        paragraph += f" доступ с {start} по {end}"
    if purpose:
        paragraph += f", с целью: {purpose}"
    paragraph += f" на объект {obj}"
    draw_text(paragraph, x=left + 8 * mm, max_width=right - left - 8 * mm)

    rooms = request_data.get("rooms") or []
    if isinstance(rooms, str):
        rooms = [part.strip() for part in rooms.split(",") if part.strip()]
    if rooms:
        draw_text("в помещения: " + ", ".join(rooms), x=left + 8 * mm, max_width=right - left - 8 * mm)

    if visitors:
        draw_text("для следующих лиц:", x=left + 8 * mm)
        for visitor in visitors:
            fio = " ".join(filter(None, [visitor.get("surname"), visitor.get("name"), visitor.get("patronymic")]))
            bits = [fio]
            if visitor.get("birth_date"):
                bits.append(f"{visitor['birth_date']}г.р.")
            series = (visitor.get("passport_series") or "").replace(" ", "")
            number = visitor.get("passport_number") or ""
            if series or number:
                bits.append(f"паспорт {series}№{number}".strip())
            if visitor.get("passport_issue_date"):
                bits.append(f"выдан {visitor['passport_issue_date']}")
            if visitor.get("passport_issued_by"):
                bits.append(visitor["passport_issued_by"])
            org_v = visitor.get("organization") or org
            line = " ".join(bits)
            if org_v:
                line = f"{line}, сотрудник {org_v}"
            draw_text(line, x=left + 10 * mm, size=9, leading=12, max_width=right - left - 10 * mm)

    if vehicles:
        draw_text("Просим разрешить въезд под погрузку/разгрузку следующим автомобилям:", x=left + 8 * mm)
        for vehicle in vehicles:
            bits = [f"гос.№ {vehicle.get('license_plate', '')}"]
            fio = " ".join(filter(None, [vehicle.get("driver_surname"), vehicle.get("driver_name"), vehicle.get("driver_patronymic")]))
            if fio:
                bits.append(fio)
            draw_text(" ".join(bits), x=left + 10 * mm, size=9)

    if equipment:
        y -= 4
        draw_text(
            "Просим разрешить внос/вынос следующего оборудования (материалов) любыми лицами, а также любыми автомобилями указанными в настоящей заявке",
            x=left + 8 * mm,
            size=9,
            max_width=right - left - 8 * mm,
        )
        cell_style = ParagraphStyle("cell", fontName=font, fontSize=7, leading=9)
        head_style = ParagraphStyle("head", fontName=font_bold, fontSize=7, leading=9)

        def cell(text, header=False):
            return Paragraph(str(text or "").replace("&", "&amp;").replace("<", "&lt;"), head_style if header else cell_style)

        header = [[cell(title, True) for title in [
            "№", "Наименование", "Марка", "Модель", "Серийный номер", "Кол.", "Ед.", "Помещение", "Ряд, место", "Внос / вынос"
        ]]]
        body = []
        for idx, item in enumerate(equipment, 1):
            body.append([
                cell(idx),
                cell(item.get("name", "")),
                cell(item.get("brand", "")),
                cell(item.get("model", "")),
                cell(item.get("serial_number", "")),
                cell(item.get("quantity", "1")),
                cell(item.get("unit", "шт.")),
                cell(item.get("room", "")),
                cell(" , ".join(part for part in (item.get("row", ""), item.get("place", "")) if part)),
                cell(item.get("action", "")),
            ])
        table = Table(header + body, colWidths=[8*mm, 22*mm, 16*mm, 20*mm, 24*mm, 10*mm, 8*mm, 30*mm, 18*mm, 18*mm])
        table.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#94A3B8")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E2E8F0")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        tw, th = table.wrap(right - left, y - bottom)
        ensure_space(th + 8)
        table.drawOn(pdf, left, y - th)
        y -= th + 10

    y -= 8
    ensure_space(50 * mm)
    director_title = request_data.get("responsible_person") or "Генеральный директор"
    director_name = request_data.get("responsible_name") or ""
    signature_y = y
    pdf.setFont(font, 11)
    pdf.drawString(left, y, director_title)
    pdf.drawRightString(right, y, director_name)
    y -= 14
    pdf.drawString(left, y, org)
    y -= 28
    pdf.setFont(font, 10)
    pdf.drawString(left + 6 * mm, y, "М.П.")
    draw_round_stamp(pdf, stamp_path, width, signature_y)
    pdf.save()
    return output_path


def suggested_filenames(request_data):
    number = re.sub(r"[^\w.\-]+", "_", request_data.get("request_number") or "zayavka")
    date = (request_data.get("work_start_date") or datetime.now().strftime("%d.%m.%Y")).replace(".", "")
    base = f"{number}_{date}" if request_data.get("request_number") else f"Заявка_{date}"
    return {
        "xlsx": OUTPUT_DIR / f"{base}.xlsx",
        "pdf": OUTPUT_DIR / f"{base}.pdf",
    }
